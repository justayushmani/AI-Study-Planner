import io
from typing import Tuple
import pypdf
import pdfplumber
import docx

class DocumentParser:
    @staticmethod
    def extract_text_from_pdf(file_bytes: bytes) -> Tuple[str, int]:
        """
        Extracts text from PDF bytes using pdfplumber with pypdf fallback.
        Returns (extracted_text, page_count).
        """
        text_parts = []
        page_count = 0

        try:
            with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                page_count = len(pdf.pages)
                for page_idx, page in enumerate(pdf.pages):
                    page_text = page.extract_text()
                    if page_text:
                        text_parts.append(f"--- Page {page_idx + 1} ---\n{page_text}")
        except Exception:
            # Fallback to PyPDF
            try:
                reader = pypdf.PdfReader(io.BytesIO(file_bytes))
                page_count = len(reader.pages)
                for page_idx, page in enumerate(reader.pages):
                    page_text = page.extract_text()
                    if page_text:
                        text_parts.append(f"--- Page {page_idx + 1} ---\n{page_text}")
            except Exception as e:
                raise ValueError(f"Failed to extract text from PDF: {str(e)}")

        full_text = "\n\n".join(text_parts).strip()
        if not full_text:
            raise ValueError(
                "No readable digital text could be extracted from this PDF. "
                "The file may be a scanned image without a text layer."
            )
        return full_text, page_count

    @staticmethod
    def extract_text_from_docx(file_bytes: bytes) -> str:
        """Extracts text from DOCX file bytes."""
        try:
            doc = docx.Document(io.BytesIO(file_bytes))
            full_text = []
            for para in doc.paragraphs:
                if para.text.strip():
                    full_text.append(para.text.strip())
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        full_text.append(row_text)
            return "\n".join(full_text).strip()
        except Exception as e:
            raise ValueError(f"Failed to extract text from DOCX document: {str(e)}")
